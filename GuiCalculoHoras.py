import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl

class Ponto:
    def __init__(self, arquivo_excel):
        self.nomeArquivo = arquivo_excel
        self.wb = openpyxl.load_workbook(self.nomeArquivo, read_only=True)
        self.ws = self.wb['Plan1']
        self.atraso = []
        self.excedente = []
        self.total_tpd = []

    def lendoTexto(self):
        self.atraso = []
        self.excedente = []
        self.total_tpd = []

        for i in range(1, 200):
            frase = str(self.ws[f'A{i}'].value or '')
            valor = str(self.ws[f'B{i}'].value or '00:00:00')
            valor = valor[:5]  # Considerar apenas HH:MM

            prefixo = frase[:3]

            if prefixo in ['021', '008']:
                self.atraso.append(valor)
            elif prefixo in ['011', '002']:
                self.excedente.append(valor)
            elif prefixo == '400':
                self.total_tpd.append(valor)

    def calcular_horas(self, saldoH, saldoM):
        def somar_tempos(lista):
            horas, minutos = 0, 0
            for tempo in lista:
                h, m = map(int, tempo.split(':'))
                horas += h
                minutos += m
            horas += minutos // 60
            minutos = minutos % 60
            return horas, minutos

        h_ex, m_ex = somar_tempos(self.excedente)
        h_neg, m_neg = somar_tempos(self.atraso)
        h_tpd, m_tpd = somar_tempos(self.total_tpd)

        total_horas = h_ex - h_neg + saldoH
        total_minutos = m_ex - m_neg + saldoM

        ajuste_h, ajuste_m = divmod(total_minutos, 60)
        total_horas += ajuste_h
        total_minutos = ajuste_m

        resultado = (
            f"SALDO ANTERIOR DE HORAS: {saldoH}h {saldoM}m\n"
            f"TOTAL DE HORAS POSITIVAS: {h_ex}h {m_ex}m\n"
            f"TOTAL DE HORAS NEGATIVAS: {h_neg}h {m_neg}m\n\n"
            f"SALDO FINAL DE HORAS: {total_horas}h {total_minutos}m\n\n"
            f"TOTAL TPD (cód. 400): {h_tpd}h {m_tpd}m\n\n"
            f"=-=-=-=-=-=-=-=-=-=-\n"
            f"Excedente: {self.excedente}\n"
            f"Atraso: {self.atraso}\n"
            f"TPD (400): {self.total_tpd}"
        )
        return resultado


class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Cálculo de Horas - Ponto")

        self.arquivo_excel = None

        tk.Label(root, text="Nome:").grid(row=0, column=0, sticky='e')
        self.nome_entry = tk.Entry(root)
        self.nome_entry.grid(row=0, column=1, columnspan=2, sticky='we')

        tk.Label(root, text="Saldo Anterior (Horas):").grid(row=1, column=0, sticky='e')
        self.horas_entry = tk.Entry(root, width=5)
        self.horas_entry.grid(row=1, column=1)

        tk.Label(root, text="Minutos:").grid(row=1, column=2, sticky='e')
        self.minutos_entry = tk.Entry(root, width=5)
        self.minutos_entry.grid(row=1, column=3)

        self.btn_arquivo = tk.Button(root, text="Selecionar Arquivo Excel", command=self.selecionar_arquivo)
        self.btn_arquivo.grid(row=2, column=0, columnspan=4, pady=5)

        self.btn_calcular = tk.Button(root, text="Calcular", command=self.calcular)
        self.btn_calcular.grid(row=3, column=0, columnspan=4, pady=5)

        self.resultado_text = tk.Text(root, height=20, width=80)
        self.resultado_text.grid(row=4, column=0, columnspan=4, pady=10)

    def selecionar_arquivo(self):
        self.arquivo_excel = filedialog.askopenfilename(
            title="Selecione o arquivo Excel",
            filetypes=[("Planilhas Excel", "*.xlsx")]
        )
        if self.arquivo_excel:
            messagebox.showinfo("Arquivo selecionado", f"Arquivo: {self.arquivo_excel}")

    def calcular(self):
        if not self.arquivo_excel:
            messagebox.showwarning("Aviso", "Por favor, selecione um arquivo Excel primeiro.")
            return

        nome = self.nome_entry.get()
        try:
            saldoH = int(self.horas_entry.get())
            saldoM = int(self.minutos_entry.get())
        except ValueError:
            messagebox.showerror("Erro", "Horas e minutos devem ser números inteiros.")
            return

        ponto = Ponto(self.arquivo_excel)
        ponto.lendoTexto()
        resultado = ponto.calcular_horas(saldoH, saldoM)

        self.resultado_text.delete("1.0", tk.END)
        self.resultado_text.insert(tk.END, f"=-=-=-=-=-=-=-=-=-=-\n{nome}\n=-=-=-=-=-=-=-=-=-=-\n\n")
        self.resultado_text.insert(tk.END, resultado)


if __name__ == '__main__':
    root = tk.Tk()
    app = App(root)
    root.mainloop()
