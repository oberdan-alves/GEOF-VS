import openpyxl

class Ponto:

    def __init__(self, nome):
        self.nome = nome
        self.nomeArquivo = 'Pasta1.xlsx'
        self.wb = openpyxl.load_workbook(self.nomeArquivo, read_only=True)
        self.ws = self.wb['Plan1']
        self.atraso = []
        self.excedente = []
        self.total_tpd = []

    def lendoTexto(self):
        self.atraso.clear()
        self.excedente.clear()
        self.total_tpd.clear()

        for i in range(1, 200):
            frase = str(self.ws[f'A{i}'].value or '')
            valor = str(self.ws[f'B{i}'].value or '00:00:00')  # agora com 3 partes

            prefixo = frase[:3]

            if prefixo in ['021', '008']:
                self.atraso.append(valor)
            elif prefixo in ['011', '002']:
                self.excedente.append(valor)
            elif prefixo == '400':
                self.total_tpd.append(valor)

    def calcular_horas(self, saldoH, saldoM):
        def somar_tempos(lista):
            horas, minutos = 0, 0
            for tempo in lista:
                try:
                    partes = tempo.split(':')
                    if len(partes) == 3:
                        h, m, _ = partes
                    elif len(partes) == 2:
                        h, m = partes
                    else:
                        continue
                    horas += int(h)
                    minutos += int(m)
                except:
                    continue
            horas += minutos // 60
            minutos = minutos % 60
            return horas, minutos

        # Soma horas
        h_ex, m_ex = somar_tempos(self.excedente)
        h_neg, m_neg = somar_tempos(self.atraso)
        h_tpd, m_tpd = somar_tempos(self.total_tpd)

        # Exibe totais parciais
        print(f"\nSALDO ANTERIOR DE HORAS: {saldoH}h {saldoM}m")
        print(f"TOTAL DE HORAS POSITIVAS: {h_ex}h {m_ex}m")
        print(f"TOTAL DE HORAS NEGATIVAS: {h_neg}h {m_neg}m")

        # Cálculo do saldo total
        total_horas = h_ex - h_neg + saldoH
        total_minutos = m_ex - m_neg + saldoM

        # Ajustar minutos
        if total_minutos < 0:
            ajuste_h, ajuste_m = divmod(-total_minutos, 60)
            total_horas -= ajuste_h
            total_minutos = -ajuste_m if ajuste_m != 0 else 0
        else:
            ajuste_h, ajuste_m = divmod(total_minutos, 60)
            total_horas += ajuste_h
            total_minutos = ajuste_m

        # Exibir saldo final
        print(f"\nSALDO FINAL DE HORAS: {total_horas}h {total_minutos}m")
        print(f"\nTOTAL TPD (cód. 400): {h_tpd}h {m_tpd}m")

        # Bloco final
        separador = '=-' * 10
        print(f"\n{separador}\n{self.nome}\n{separador}")
        print(f"Horas: {total_horas}")
        print(f"Min: {total_minutos}")
        print(f"Total TPD: {h_tpd}h {m_tpd}m")
        print(f"Excedente: {self.excedente}")
        print(f"Atraso: {self.atraso}")
        print(f"TPD (400): {self.total_tpd}")




if __name__ == '__main__':
    oberdan = Ponto('Oberdan')
    oberdan.lendoTexto()
    oberdan.calcular_horas(-16, -53)
