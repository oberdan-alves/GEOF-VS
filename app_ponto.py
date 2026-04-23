import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
import io
import sys

class Ponto:
    def __init__(self, arquivo_excel):
        self.nomeArquivo = arquivo_excel
        self.wb = openpyxl.load_workbook(self.nomeArquivo, read_only=True)
        self.ws = self.wb['Plan1']
        self.atraso = []
        self.excedente = []
        self.total_tpd = []

    def lendoTexto(self):
        self.atraso.clear()
        self.excedente.clear()
        self.total_tpd.clear()

        for i in range(1, 200):
            frase = str(self.ws[f'A{i}'].value or '')
            valor = str(self.ws[f'B{i}'].value or '00:00:00')

            prefixo = frase[:3]

            if prefixo in ['021', '008']:
                self.atraso.append(valor)
            elif prefixo in ['011', '002']:
                self.excedente.append(valor)
            elif prefixo == '400':
                self.total_tpd.append(valor)

    def calcular_horas(self, saldoH, saldoM):
        def somar_tempos(lista):
            horas, minutos = 0, 0
            for tempo in lista:
                try:
                    partes = tempo.split(':')
                    if len(partes) == 3:
                        h, m, _ = partes
                    elif len(partes) == 2:
                        h, m = partes
                    else:
                        continue
                    horas += int(h)
                    minutos += int(m)
                except:
                    continue
            horas += minutos // 60
            minutos = minutos % 60
            return horas, minutos

        h_ex, m_ex = somar_tempos(self.excedente)
        h_neg, m_neg = somar_tempos(self.atraso)
        h_tpd, m_tpd = somar_tempos(self.total_tpd)

        print(f"\nSALDO ANTERIOR DE HORAS: {saldoH}h {saldoM}m")
        print(f"TOTAL DE HORAS POSITIVAS: {h_ex}h {m_ex}m")
        print(f"TOTAL DE HORAS NEGATIVAS: {h_neg}h {m_neg}m")

        total_horas = h_ex - h_neg + saldoH
        total_minutos = m_ex - m_neg + saldoM

        if total_minutos < 0:
            ajuste_h, ajuste_m = divmod(-total_minutos, 60)
            total_horas -= ajuste_h
            total_minutos = -ajuste_m if ajuste_m != 0 else 0
        else:
            ajuste_h, ajuste_m = divmod(total_minutos, 60)
            total_horas += ajuste_h
            total_minutos = ajuste_m

        separador = '=-' * 10
        print(f"\nSALDO FINAL DE HORAS: {total_horas}h {total_minutos}m")
        print(f"\nTOTAL TPD (cód. 400): {h_tpd}h {m_tpd}m")
        print(f"\n{separador}")
        print(f"Horas: {total_horas}")
        print(f"Min: {total_minutos}")
        print(f"Total TPD: {h_tpd}h {m_tpd}m")
        print(f"Excedente: {self.excedente}")
        print(f"Atraso: {self.atraso}")
        print(f"TPD (400): {self.total_tpd}")

class AppPonto:
    def __init__(self, root):
        self.root = root
        self.root.title("Cálculo de Ponto")
        self.root.geometry("700x500")
        self.root.configure(bg="#f0f0f0")

        self.arquivo_excel = None

        font_title = ("Arial", 20, "bold")
        font_label = ("Arial", 14)
        font_button = ("Arial", 14)

        tk.Label(root, text="Cálculo de Horas", font=font_title, bg="#f0f0f0").pack(pady=10)

        btn_excel = tk.Button(root, text="Selecionar Arquivo Excel", font=font_button, command=self.selecionar_excel)
        btn_excel.pack(pady=5)

        frame = tk.Frame(root, bg="#f0f0f0")
        frame.pack(pady=10)

        tk.Label(frame, text="Saldo Horas:", font=font_label, bg="#f0f0f0").grid(row=0, column=0, padx=5, pady=5)
        self.entry_h = tk.Entry(frame, font=font_label, width=5)
        self.entry_h.grid(row=0, column=1)

        tk.Label(frame, text="Saldo Minutos:", font=font_label, bg="#f0f0f0").grid(row=0, column=2, padx=5, pady=5)
        self.entry_m = tk.Entry(frame, font=font_label, width=5)
        self.entry_m.grid(row=0, column=3)

        self.output_text = tk.Text(root, font=("Courier", 12), width=80, height=15)
        self.output_text.pack(pady=10)

        btn_calcular = tk.Button(root, text="Calcular", font=font_button, command=self.calcular)
        btn_calcular.pack(pady=10)

    def selecionar_excel(self):
        caminho = filedialog.askopenfilename(filetypes=[("Planilhas Excel", "*.xlsx")])
        if caminho:
            self.arquivo_excel = caminho
            messagebox.showinfo("Arquivo selecionado", f"Arquivo carregado:\n{caminho}")

    def calcular(self):
        if not self.arquivo_excel:
            messagebox.showwarning("Atenção", "Selecione um arquivo Excel primeiro.")
            return

        try:
            h = int(self.entry_h.get())
            m = int(self.entry_m.get())
        except ValueError:
            messagebox.showerror("Erro", "Informe valores numéricos para horas e minutos.")
            return

        try:
            ponto = Ponto(self.arquivo_excel)
            ponto.lendoTexto()

            old_stdout = sys.stdout
            sys.stdout = buffer = io.StringIO()

            ponto.calcular_horas(h, m)

            sys.stdout = old_stdout
            resultado = buffer.getvalue()
            self.output_text.delete("1.0", tk.END)
            self.output_text.insert(tk.END, resultado)

        except Exception as e:
            messagebox.showerror("Erro", f"Ocorreu um erro ao processar o arquivo:\n{e}")

if __name__ == '__main__':
    root = tk.Tk()
    app = AppPonto(root)
    root.mainloop()
    'pyinstaller --onefile --noconsole app_ponto.py'
