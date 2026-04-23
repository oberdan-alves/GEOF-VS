import openpyxl

class Ponto:

    def __init__(self, nome):
        self.nome = nome
        self.nomeArquivo = 'Pasta1.xlsx'
        self.wb = openpyxl.load_workbook(self.nomeArquivo, read_only=True)
        self.ws = self.wb['Plan1']
        self.atraso = []
        self.excedente = []

    def lendoTexto(self):
        for i in range(1, 200):
            frase = str(self.ws['A' + str(i)].value)
            if frase[0:3] == '021':
                self.atraso.append(str(self.ws['B' + str(i)].value))
            if frase[0:3] == '011':
                self.excedente.append(str(self.ws['B' + str(i)].value))

    def calcular_horas(self, saldoH, saldoM):

        horasex = []
        horasneg = []

        for i in self.excedente:
            eita = i.split(':')
            horasex.append(eita)

        for b in self.atraso:
            eita = b.split(':')
            horasneg.append(eita)

        print(f'horas excedentes: {horasex}')
        print(f'horas negativas: {horasneg}')
        

        hpos = 0
        hneg = 0
        mpos = 0
        mneg = 0

        for x in range(0, len(self.excedente)):
            hpos += int(horasex[x][0])
            mpos += int(horasex[x][1])

        for x in range(0, len(self.atraso)):
            hneg += int(horasneg[x][0])
            mneg += int(horasneg[x][1])

        print(f'TOTAL DE HORAS POSITIVAS: {hpos}')
        print(f'TOTAL DE MINUTOS POSITIVOS: {mpos}')
        minpos = divmod(mpos,60)
        print(f'divimod dos minutos {minpos}')
        horaspositivas = hpos + minpos[0]
        minutpositivas = minpos[1]
        print(f'horas posivitvas calculadas: {horaspositivas}')
        print(f'minutos positivas calculadas: {minutpositivas}')
        print()
        print(f'TOTAL DE HORAS NEGATIVAS: {hneg}')
        print(f'TOTAL DE MINUTOS NEGATIVOS: {mneg}')
        minneg = divmod(mneg, 60)
        print(f'divimod dos minutos {minneg}')
        horasnegativas = hneg + minneg[0]
        minutosnegativos = minneg[1]
        print(f'horas negativas calculadas: {horasnegativas}')
        print(f'minutos negativas calculadas: {minutosnegativos}')
        print()

        HORAS_TOTAIS = horaspositivas - horasnegativas + saldoH
        print(f'HORAS: {HORAS_TOTAIS}')
        
        MINUTOS_TOTAIS = minutpositivas - minutosnegativos + saldoM
        print(f'MINUTOS TOTAIS: {MINUTOS_TOTAIS}')
        if MINUTOS_TOTAIS>=0:
            hor = divmod(MINUTOS_TOTAIS, 60)
        else:
            hor = divmod(MINUTOS_TOTAIS, -60)


        print(f'{hor}\n')
        HORAS_TOTAIS += hor[0]
        a = '=-'*10
        print(f'{a}\n{self.nome}\n{a}\nHoras: {HORAS_TOTAIS}\nMin: {hor[1]}\n')
        print(self.excedente)
        print(self.atraso)


if __name__ == '__main__':
    oberdan = Ponto('Oberdan')
    oberdan.lendoTexto()
    #print(f'Excedente = {oberdan.excedente}\n'
    #      f'Atraso = {oberdan.atraso}')
    #oberdan.calcular_horas(-5, 1)
